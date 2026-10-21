"""Regression checks for the Oct 20 2026 quote detail passes and their proposal rendering.

Run from the repo root:  python3 tests/quote_details_test.py
Self-contained: synthetic quote text in the layouts pdftotext -layout produces, no OpenAI calls.

  * terrorism_liability.py - the liability extension schedule of a terrorism policy
    (AEGIS Terror Shield style: "Limit of Liability: USD 2,000,000 each and every Claim ...")
  * workers_comp_quote.py  - a Markel / mPortal style WC quote (Payroll, Net Rate, PR Net Rate,
    Rate, Premium; premium build-up; payment plans; an out-of-state fraud notice)
  * program_sponsor.py     - a member umbrella quote whose "Named Insured" is the program sponsor
  * proposal_generator.py  - "Premium Indication", the Terrorism Liability Extension table, the
    WC Premium Development table, admitted WC, no sponsor among the named insureds
"""
import copy
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import program_sponsor as ps  # noqa: E402
import terrorism_liability as tl  # noqa: E402
import workers_comp_quote as wq  # noqa: E402

PASS = 0


def check(cond, msg):
    global PASS
    if not cond:
        raise AssertionError(msg)
    PASS += 1


TERROR_TEXT = """
--- Page 3 ---
6.   Sum Insured:                                USD 8,000,000.00 each Loss Occurrence and in the aggregate during
                                                 USD 8,000,000.00 the Policy Period.
     B. Active Assailant Event:                  USD 1,000,000            each Loss Occurrence and in the aggregate

--- Page 41 ---
                           General Liability Extension (Act of Terrorism, Act of
ENDORSEMENT
                           Sabotage) (Claims Made and Reported)


                                                        SCHEDULE

1.   Named Insured:                                As stated in the Schedule to the Policy to which this Extension attaches

3.   Limit of Liability:                           USD 2,000,000 each and every Claim and in the aggregate
                                                   USD 2,000,000 during the Policy Period.
                                                   Such Limit of Liability is part of and not in addition to the Sum Insured set
                                                   forth in the Schedule of the Policy to which this Extension forms a part.
                                                   Defense, Settlement of Claims, and Supplemental Payments are payable within the Limit of
                                                   Liability and not in addition to such Limit.

4.   Deductible:                                   USD 5,000 each and every Claim.

5.   Nominee for Notice of Claims or               As stated in the Schedule to the Policy to which this Extension attaches
     Circumstances to:

6.   Retroactive Date:                             30 October 2026

                                                          [ 40 ]

--- Page 42 ---
 SECTION I            Insuring Agreements

NOTICE TO THE INSURED - This is a claims made and reported Policy.
The Underwriters do not have a duty to defend.
The Limit of Liability of this Extension available to pay damages may be impaired
by the reduction or exhaustion of the aggregate Sum Insured of the Policy that this Extension attaches.
       a.   an Act of Terrorism or an Act of Sabotage; or
      practicable and in no event later than ninety (90) days after the expiration of the Policy Period and in
                                                      [ 41 ]        Terrorism Liability Extension to AEGIS Terror Shield
"""

WC_TEXT = """
                                                                              Client Information
                                                                              Company          Sample Hotel LLC
                                                                              Effective Date   12/01/2026
Quote Proposal (Expires 01/15/2027)
Thank you for giving us the opportunity to provide the following workers' compensation quote for the captioned insured with Markel
Insurance Company. The quote is as follows:
                                                                                               Net         PR Net
State Code                           Description                               Payroll                                  Rate Premium
                                                                                               Rate         Rate
               Hotel - All Other Employees & Salespersons,
GA      9052                                                                $500,000.00 0.910*          1.010           1.200 $6,000
               Drivers
                                                                            $500,000                                            $6,000
Liability Limit:                                                                            1000/1000/1000                      $250
GA - Schedule Rating                                                                        -0.250                              -$1,563
GA E-MOD:                                                                                   0.950                               -$234
Premium Discount:                                                                                                               $0
Expense Constant:                                                                                                               $160
Terrorism Charge:                                                                                                               $50
Final Premium                                                                                                                   $4,663.00
Payment Plan Options
                                  Down
           Payment Plan                  Down Payment $
                               Payment %
1-Pay                          100%      $4,663.00
4-Pay                          40%       $1,865.00
PayGo, Payroll Reporting       0%        $0.00

Re: Quote Number 123
12/01/2026 - 12/01/2027
                                       Markel Insurance Company
                                  CALIFORNIA FRAUD WARNING
"""

UMB_TEXT = """
Great Point Insurance is pleased to provide you with our quotation for Sample Hotel LLC offered through the Paramount Real Estate
Group, Inc. Members of the Paramount Real Estate Group, Inc. may not be protected by the insurance insolvency guarantee fund.
Named Insured:                             Paramount Real Estate Group, Inc.
Insured:                                   Sample Hotel LLC
Member Coverage Period:                    From: 12/01/2026 To: 12/01/2027
"""

# ---------------------------------------------------------------- terrorism liability
ext = tl.find_liability_extension(TERROR_TEXT)
check(ext is not None, "extension found")
check(ext["limit_each"] == 2000000 and ext["limit_aggregate"] == 2000000, f"limits {ext}")
check(ext["deductible_amount"] == 5000, "deductible")
check(ext["retroactive_date"] == "10/30/2026", "retro date")
check(ext["trigger"] == "Claims made and reported", "trigger")
check(ext["reporting_days"] == 90, "reporting window")
check(ext["defense_within_limits"] and ext["no_duty_to_defend"] and ext["part_of_sum_insured"], "flags")
check(ext["source_page"] == 41 and ext["printed_page"] == 40, "page numbers")
check(ext["title"].startswith("General Liability Extension (Act of Terrorism, Act of Sabotage)"), ext["title"])
check(tl.find_liability_extension("SECTION IV Extensions ... Limit of Liability (prose)") is None, "no schedule, no extension")

data = {"coverages": {"terrorism": {"carrier": "Certain Underwriters at Lloyd's, London",
                                    "limits": [{"description": "General Liability Extension", "limit": "per policy"}]},
                      "general_liability": {"carrier": "Southlake Specialty Insurance Company",
                                            "limits": [{"description": "Each Occurrence", "limit": "$1,000,000"}]}}}
notes = tl.run_terrorism_liability_pass(data, [{"filename": "terror.pdf", "text": TERROR_TEXT}])
tc = data["coverages"]["terrorism"]
check(len(notes) == 1 and "printed page 40" in notes[0], notes)
check(tc["liability_extension"]["limit_each"] == 2000000, "attached to terrorism")
rows = [r for r in tc["limits"] if "Liability" in r["description"]]
check(len(rows) == 1 and rows[0]["limit"].startswith("$2,000,000 each and every Claim"), rows)
check(data["coverages"]["general_liability"]["limits"] == [{"description": "Each Occurrence", "limit": "$1,000,000"}],
      "GL untouched")
check(tl.run_terrorism_liability_pass({"coverages": {}}, [{"filename": "t", "text": TERROR_TEXT}]) == [],
      "no terrorism coverage, nothing to do")

# ---------------------------------------------------------------- workers comp
p = wq.parse_wc_quote(WC_TEXT)
check(p is not None and p["carrier"] == "Markel Insurance Company", p and p["carrier"])
r = p["class_rows"][0]
check(r["rate"] == "1.200" and r["premium"] == "$6,000", f"manual rate picked by arithmetic: {r}")
check(r["classification"] == "Hotel - All Other Employees & Salespersons, Drivers", r["classification"])
check(p["final_premium"] == 4663 and p["el_limits"] == [1000000] * 3, "final premium / EL")
check([pl["plan"] for pl in p["payment_plans"]] == ["1-Pay", "4-Pay", "PayGo, Payroll Reporting"], p["payment_plans"])
check(p["policy_period"] == "12/01/2026 - 12/01/2027", p["policy_period"])
dev = wq.premium_development(p)
check(dev[0]["item"] == "Manual Premium ($500,000 payroll x 1.200)" and dev[-1]["item"] == "Final Premium", dev)
check(any(d["factor"] == "$1M / $1M / $1M" for d in dev), dev)
check(any(d["item"] == "Schedule Rating (GA)" and d["factor"] == "-25%" for d in dev), dev)
check(any(d["item"] == "Experience Modification (GA)" and d["factor"] == "0.950" for d in dev), dev)

wcdata = {"coverages": {"workers_comp": {
    "carrier": "Markel Insurance Company", "carrier_admitted": False, "premium": 4663, "total_premium": 4663,
    "limits": [{"description": "EL - Each Accident", "limit": "$1,000,000"}],
    "rating_basis": [{"state": "GA", "class_code": "9052", "payroll": "$500,000", "rate": "0.910"}],
    "forms_endorsements": [{"form_number": "MPIL 1157-CA 05 23",
                            "description": "California Fraud Warning Advisory Notice to Policyholders"},
                           {"form_number": "WC 10 00 00", "description": "Georgia Amendatory Endorsement"}]}},
    "payment_options": [{"carrier": "Markel Insurance Company", "coverage_type": "Workers Compensation",
                         "terms": "1-Pay; Down Payment 100%"},
                        {"carrier": "Markel Insurance Company", "coverage_type": "Workers Compensation",
                         "terms": "4-Pay; Down Payment 40%"},
                        {"carrier": "Starr Surplus Lines Insurance Company", "coverage_type": "Property",
                         "terms": "Due in 30 days"}]}
wq.run_workers_comp_pass(wcdata, [{"filename": "wc.pdf", "text": WC_TEXT}])
w = wcdata["coverages"]["workers_comp"]
check(w["carrier_admitted"] is True, "WC admitted")
check(w["limits"][0]["limit"] == "Statutory (GA)", w["limits"])
check(len([x for x in w["limits"] if x["description"].startswith("EL -")]) == 3, w["limits"])
check(w["rating_basis"][0]["rate"] == "1.200" and w["rating_basis"][0]["premium"] == "$6,000", w["rating_basis"])
check([f["form_number"] for f in w["forms_endorsements"]] == ["WC 10 00 00"], w["forms_endorsements"])
wc_pay = [po for po in wcdata["payment_options"] if "Markel" in po["carrier"]]
check(len(wc_pay) == 1 and "PayGo, Payroll Reporting: no down payment" in wc_pay[0]["terms"], wc_pay)
check(any("Starr" in po["carrier"] for po in wcdata["payment_options"]), "other carriers kept")
sl = {"coverages": {"workers_comp": {"carrier": "Sample Mutual", "carrier_admitted": False}}}
wq.run_workers_comp_pass(sl, [{"filename": "x", "text": "workers' compensation ... surplus lines ... no table"}])
check(sl["coverages"]["workers_comp"]["carrier_admitted"] is True, "WC admitted even when unparsed")

# ---------------------------------------------------------------- program sponsor
found = ps.detect_program_sponsors([{"filename": "umb.pdf", "text": UMB_TEXT}])
check(any(ps._norm(k) == ps._norm("Paramount Real Estate Group, Inc.") for k in found), found)
check(not ps.is_program_sponsor("Sunrise Hotel Group LLC", ["Hotel Group, Inc."]), "no substring matches")
check(not ps.is_program_sponsor("Paramount Hospitality LLC"), "client named Paramount is kept")
sp = {"client_info": {"named_insured": "Sample Hotel LLC"},
      "named_insureds": [{"name": "Sample Hotel LLC"}, {"name": "PARAMOUNT REAL ESTATE GROUP, INC."}],
      "coverages": {"umbrella": {"carrier": "StarStone National Insurance Company",
                                 "named_insured": "Paramount Real Estate Group, Inc.",
                                 "additional_named_insureds": [{"name": "Paramount Real Estate Group, Inc."}]}}}
notes = ps.run_program_sponsor_pass(sp, [{"filename": "umb.pdf", "text": UMB_TEXT}])
check([n["name"] for n in sp["named_insureds"]] == ["Sample Hotel LLC"], sp["named_insureds"])
check(sp["coverages"]["umbrella"]["named_insured"] == "Sample Hotel LLC", sp["coverages"]["umbrella"])
check(sp["coverages"]["umbrella"]["additional_named_insureds"] == [], "umbrella ANI scrubbed")
check(sp["program_sponsors"] and notes, "sponsor recorded")

# ---------------------------------------------------------------- page scoring
import proposal_extractor as pe  # noqa: E402
sched_page = TERROR_TEXT.split("--- Page 41 ---")[1].split("--- Page 42 ---")[0]
check(pe._score_page(sched_page) >= 50, f"schedule page is critical: {pe._score_page(sched_page)}")

# ---------------------------------------------------------------- generator
import docx  # noqa: E402
import proposal_generator as pg  # noqa: E402

gen = {
    "client_name": "Sample Hotel LLC",
    "client_info": {"named_insured": "Sample Hotel LLC", "address": "1 Main St, Atlanta, GA 30301",
                    "effective_date": "12/01/2026", "expiration_date": "12/01/2027"},
    "named_insureds": [{"name": "Sample Hotel LLC"}, {"name": "Paramount Real Estate Group, Inc."}],
    "program_sponsors": [],
    "locations": [{"number": "1", "name": "Sample Inn", "address": "1 Main St", "city": "Atlanta", "state": "GA",
                   "zip": "30301", "tiv": 8000000}],
    "coverages": {
        "workers_comp": copy.deepcopy(w),
        "terrorism": dict(copy.deepcopy(tc), carrier="Certain Underwriters at Lloyd's, London", premium=2500,
                          total_premium=2500, carrier_admitted=False,
                          limits=[{"description": "Act of Terrorism - each Loss Occurrence", "limit": "USD 8,000,000.00"}]
                          + [x for x in tc["limits"] if "Liability" in x["description"]]),
        "umbrella": {"carrier": "StarStone National Insurance Company", "premium": 5000, "total_premium": 5000,
                     "named_insured": "Paramount Real Estate Group, Inc.",
                     "limits": [{"description": "Each Occurrence", "limit": "$5,000,000"}]},
        "cyber": {"carrier": "", "premium": 0, "optional": True},
    },
    "payment_options": wcdata["payment_options"],
}
with tempfile.TemporaryDirectory() as td:
    out = os.path.join(td, "p.docx")
    pg._sanitize_for_render(gen)
    pg.generate_proposal(gen, out)
    d = docx.Document(out)
    text = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
check("Premium Indication" in cells, "optional coverages column")
check("Terrorism Liability Extension" in text and "Claims made and reported" in cells, "terrorism extension table")
check("10/30/2026" in cells and "$8,000,000" in cells and "Usd" not in cells and "USD 8" not in cells, "terrorism values")
check("Premium Development" in text and "Final Premium" in cells, "WC premium development")
check("Paramount Real Estate Group" not in cells, "sponsor never listed")
wc_rows = [r for t in d.tables for r in t.rows if r.cells[0].text == "Admitted Status"]
check(wc_rows and wc_rows[0].cells[1].text == "Admitted", "WC section admitted")
check("Georgia Amendatory Endorsement" in cells and "California Fraud Warning" not in cells, "WC forms")

gen2 = copy.deepcopy(gen)
gen2["coverages"]["workers_comp"]["forms_endorsements"] = []
with tempfile.TemporaryDirectory() as td:
    out = os.path.join(td, "p2.docx")
    pg.generate_proposal(gen2, out)
    text2 = "\n".join(p.text for p in docx.Document(out).paragraphs)
check("filed workers compensation and employers liability policy forms for Georgia" in text2, "WC forms note")

print(f"quote_details_test: {PASS} checks passed")
