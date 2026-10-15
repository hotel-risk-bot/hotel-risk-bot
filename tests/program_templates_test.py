"""Regression checks for program_templates.py (Starr property / PartnerRe umbrella templates).

Run from the repo root:  python3 tests/program_templates_test.py
Self-contained: synthetic SOV and GL data, no OpenAI calls. Renders a proposal through
generate_proposal() and checks the text of the .docx.
"""
import copy
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from program_templates import apply_program_templates, is_tier1_county  # noqa: E402

PASS = 0


def check(cond, msg):
    global PASS
    if not cond:
        raise AssertionError(msg)
    PASS += 1


def sov(locs):
    rows = []
    for i, (dba, addr, city, st, county, tiv, aop) in enumerate(locs, 1):
        rows.append({"location_num": i, "dba": dba, "corporate_name": f"{dba} LLC", "address": addr,
                     "city": city, "state": st, "zip_code": "77550", "county": county,
                     "building_value": tiv * 0.8, "contents_value": tiv * 0.07, "bi_value": tiv * 0.13,
                     "other_value": 0, "tiv": tiv, "aop_deductible": aop})
    return {"summary": {"named_insured": "Gulf Hotels LLC", "mailing_address": "1 Seawall Blvd, Galveston, TX 77550",
                        "effective_date": "2026-11-01 00:00:00"},
            "locations": [{k: v for k, v in r.items() if k != "aop_deductible"} for r in rows],
            "locations_raw": rows,
            "totals": {"tiv": sum(r["tiv"] for r in rows), "building_value": sum(r["building_value"] for r in rows),
                       "contents_value": sum(r["contents_value"] for r in rows), "bi_value": sum(r["bi_value"] for r in rows)}}


GL = {"carrier": "Berkley Specialty Insurance Company", "premium": 9000, "total_premium": 9800,
      "limits": [{"description": "Each Occurrence", "limit": "$1,000,000"},
                 {"description": "General Aggregate", "limit": "$2,000,000"},
                 {"description": "Products/Completed Operations Aggregate", "limit": "$2,000,000"},
                 {"description": "Personal & Advertising Injury", "limit": "$1,000,000"}]}

# 1. Tier 1 county detection
check(is_tier1_county("TX", "Galveston"), "Galveston TX is Tier 1")
check(is_tier1_county("FL", "Orange"), "Florida is whole-state Tier 1")
check(not is_tier1_county("TX", "Hood"), "Hood TX is not Tier 1")

# 2. Mixed portfolio: Named Storm only on the Tier 1 location, wind % with per-location amounts
S = sov([("Seawall Inn", "1 Seawall Blvd", "Galveston", "TX", "Galveston", 10_000_000, 25000),
         ("Hill Country Inn", "5 Main St", "Kerrville", "TX", "Kerr", 6_000_000, 10000)])
d = {"client_info": {}, "coverages": {"general_liability": copy.deepcopy(GL)}}
notes = apply_program_templates(d, {
    "property": {"enabled": True, "premium": 40000, "taxes_fees": 3000, "include_bm": "auto",
                 "deductibles": {"wind": {"mode": "pct", "pct": 2, "minimum": 50000},
                                 "named_storm": {"apply": "auto", "pct": 5, "minimum": 100000}}},
    "umbrella": {"enabled": True, "limit": 1000000, "premium": 3000, "total_premium": 3300}}, S)
p, u = d["coverages"]["property"], d["coverages"]["umbrella"]
deds = {x["description"]: x["amount"] for x in p["deductibles"]}
check(deds["All Other Perils (property damage and time element combined)"] == "$25,000 Per Occurrence",
      "AOP takes the highest SOV AOP deductible")
ns = [k for k in deds if k.startswith("Named Storm")]
check(len(ns) == 1 and "Seawall Inn - Galveston" in ns[0] and "Hill Country" not in ns[0],
      f"Named Storm applies to the Tier 1 location only: {ns}")
check("$500,000" in deds[ns[0]], "5% of $10M = $500,000 shown for the Tier 1 location")
wind = [k for k in deds if k.startswith("Wind")][0]
check("other than Named Storm" in wind, "wind label excludes named storm when NS applies")
check("$200,000" in deds[wind] and "$120,000" in deds[wind], "2% wind amounts per location")
check(p["total_premium"] == 43000 and p["tria_premium"] == 2000, "total = premium + taxes; TRIA 5%")
check(any(f["form_number"].startswith("PR 012") for f in p["forms_endorsements"]),
      "B&M endorsement kept when no EB quote")
check(any(a["description"].startswith("Boiler and Machinery") for a in p["additional_coverages"]), "B&M sublimits kept")
check(len(p["forms_endorsements"]) == 67, f"67 Starr forms with B&M ({len(p['forms_endorsements'])})")
check(p["effective_date"] == "11/01/2026" and p["expiration_date"] == "11/01/2027", "policy period from SOV")
check(d["client_info"]["named_insured"] == "Gulf Hotels LLC", "named insured filled from SOV")
check(u["layer_description"] == "$1M xs Primary" and u["limits"][0]["limit"] == "$1,000,000", "umbrella $1M")
check(u["underlying_insurance"][0]["carrier"] == "Berkley Specialty Insurance Company", "underlying from GL")
check(len(u["underlying_insurance"]) == 1, "no HNOA row when the GL has no HNOA")
check(u["tria_premium"] == 60, "umbrella TRIA 2%")
prem = [f for f in u["forms_endorsements"] if "0329" in f["form_number"]][0]
check("1 Seawall Blvd, Galveston, TX 77550" in prem["description"] and "5 Main St" in prem["description"],
      "designated premises lists every SOV address")
check(any(f["form_number"] == "PRE ALT SLN TX 0126" for f in u["forms_endorsements"]), "TX SL notice")
check(any("Tier 1 wind county" in n for n in notes), "Tier 1 note")
for cov in (p, u):
    txt = repr(cov).lower()
    check("amwins" not in txt and "alta" not in txt and "heartland" not in txt, "no wholesaler/MGA names")

# 3. EB quote present -> B&M auto-excluded; explicit True keeps it with a note
d2 = {"client_info": {}, "coverages": {"equipment_breakdown": {"carrier": "Travelers", "premium": 1000}}}
apply_program_templates(d2, {"property": {"enabled": True, "premium": 1}}, S)
p2 = d2["coverages"]["property"]
check(not any(f["form_number"].startswith("PR 012") for f in p2["forms_endorsements"]), "B&M form dropped")
check(not any(a["description"].startswith("Boiler") for a in p2["additional_coverages"]), "B&M sublimits dropped")
check(any(r["label"] == "Boiler & Machinery" and "separate policy" in r["value"] for r in p2["summary_rows"]),
      "B&M summary row says separate policy")
d3 = {"client_info": {}, "coverages": {"equipment_breakdown": {"carrier": "Travelers", "premium": 1000}}}
n3 = apply_program_templates(d3, {"property": {"enabled": True, "premium": 1, "include_bm": True}}, S)
check(any(f["form_number"].startswith("PR 012") for f in d3["coverages"]["property"]["forms_endorsements"]),
      "explicit include_bm keeps B&M")
check(any("Equipment Breakdown quote is in this proposal" in n for n in n3), "both-B&M-and-EB note")

# 4. Umbrella without a GL quote, GPT-6 nulls elsewhere, no SOV
d4 = {"client_info": None, "coverages": {"general_liability": {"carrier": None, "limits": None}},
      "locations": [{"address": "9 Elm St", "city": "Austin", "state": "TX", "zip": "78701"}]}
n4 = apply_program_templates(d4, {"umbrella": {"enabled": True, "limit": "2000000", "premium": "2,500"}}, None)
u4 = d4["coverages"]["umbrella"]
check(u4["underlying_insurance"][0]["carrier"] == "To be confirmed", "no-GL underlying placeholder")
check(any("REVIEW REQUIRED" in n and "General Liability" in n for n in n4), "no-GL review note")
check(u4["premium"] == 2500.0 and u4["total_premium"] == 2500.0, "premium parses '2,500'")
check("9 Elm St, Austin, TX 78701" in [f for f in u4["forms_endorsements"] if "0329" in f["form_number"]][0]["description"],
      "premises fall back to data.locations")

# 5. Re-apply with property unticked removes only the template property; extracted property replaced with note
d5 = {"client_info": {}, "coverages": {"property": {"carrier": "Kinsale Insurance Company", "premium": 5}}}
n5 = apply_program_templates(d5, {"property": {"enabled": True, "premium": 10}}, S)
check(d5["coverages"]["property"]["carrier"] == "Starr Surplus Lines Insurance Company", "template replaces extracted")
check(any("replaced the extracted property quote (Kinsale" in n for n in n5), "replacement note")
apply_program_templates(d5, {"property": {"enabled": False}}, S)
check("property" not in d5["coverages"], "unticking removes the template property")
check(not any("Starr" in w for w in d5["warnings"]), "stale template notes cleared")

# 6. Merged SOV: property values come from _property_sov (Amwins), not the Hospitality SOV
S6 = {"locations": [{"dba": "A", "address": "1 A St", "city": "X", "state": "GA", "tiv": 0}],
      "totals": {"tiv": 0}, "_property_sov": S}
d6 = {"client_info": {}, "coverages": {}}
apply_program_templates(d6, {"property": {"enabled": True, "premium": 1}}, S6)
check(d6["coverages"]["property"]["tiv"] == "$16,000,000", "TIV from the property SOV")

# 7. Render through the generator
from proposal_generator import generate_proposal  # noqa: E402
from docx import Document  # noqa: E402

out = os.path.join(tempfile.mkdtemp(), "t.docx")
data = copy.deepcopy(d)
data["sov_data"] = S
generate_proposal(data, out)
doc = Document(out)
text = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        text += "\n" + " | ".join(c.text for c in row.cells)
for needle in ("Starr Surplus Lines Insurance Company", "PartnerRe Insurance Solutions Bermuda Ltd.",
               "Terms & Conditions", "Named Storm, Tier 1 Wind Counties", "PRE ALT SLN TX 0126",
               "Sublimited Coverages Endorsement", "A+ (XV)", "Carrier Participation"):
    check(needle.lower() in text.lower(), f"docx contains {needle!r}")
check("Earth Movement | $25,000 Per Occurrence" in text, "EM deductible not dropped by the NOT COVERED filter")
check("amwins" not in text.lower() and "alta risk" not in text.lower(), "docx free of wholesaler names")
print(f"program_templates_test: {PASS} checks passed")
